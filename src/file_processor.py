import os
import PyPDF2
import docx
import openpyxl
from pptx import Presentation
import requests
from bs4 import BeautifulSoup
from io import BytesIO
from crewai_tools import FirecrawlScrapeWebsiteTool

class FileProcessor:
    """Process various file types and extract text."""
    @staticmethod
    def process_pdf(file_bytes):
        with BytesIO(file_bytes) as f:
            reader = PyPDF2.PdfReader(f)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text

    @staticmethod
    def process_docx(file_bytes):
        with BytesIO(file_bytes) as f:
            doc = docx.Document(f)
            return "\n".join([paragraph.text for paragraph in doc.paragraphs])

    @staticmethod
    def process_xlsx(file_bytes):
        with BytesIO(file_bytes) as f:
            wb = openpyxl.load_workbook(f)
            text = ""
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text += f"Sheet: {sheet_name}\n"
                for row in sheet.iter_rows(values_only=True):
                    text += " | ".join([str(cell) for cell in row if cell is not None]) + "\n"
            return text

    @staticmethod
    def process_pptx(file_bytes):
        with BytesIO(file_bytes) as f:
            prs = Presentation(f)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            return text

    @staticmethod
    def process_txt(file_bytes):
        return file_bytes.decode('utf-8')

    @staticmethod
    def process_url(url):
        try:
            if os.environ.get("FIRECRAWL_API_KEY"):
                scraper = FirecrawlScrapeWebsiteTool(
                    api_key=os.environ.get("FIRECRAWL_API_KEY"),
                    url=url,
                    page_options={"onlyMainContent": True}
                )
                result = scraper.run()
                return result
            else:
                # Fallback to basic request
                response = requests.get(url)
                soup = BeautifulSoup(response.content, 'html.parser')
                return soup.get_text()
        except Exception as e:
            print(f"Error scraping URL: {e}")
            # Fallback to basic request
            response = requests.get(url)
            soup = BeautifulSoup(response.content, 'html.parser')
            return soup.get_text()

    @classmethod
    def process_file(cls, file, file_type=None):
        if file_type is None:
            file_name = file.name.lower()
            if file_name.endswith('.pdf'):
                return cls.process_pdf(file.read())
            elif file_name.endswith('.docx'):
                return cls.process_docx(file.read())
            elif file_name.endswith('.xlsx') or file_name.endswith('.xls'):
                return cls.process_xlsx(file.read())
            elif file_name.endswith('.pptx') or file_name.endswith('.ppt'):
                return cls.process_pptx(file.read())
            elif file_name.endswith('.txt'):
                return cls.process_txt(file.read())
            else:
                return "Unsupported file format"
        else:
            if file_type == 'pdf':
                return cls.process_pdf(file.read())
            elif file_type == 'docx':
                return cls.process_docx(file.read())
            elif file_type in ['xlsx', 'xls']:
                return cls.process_xlsx(file.read())
            elif file_type in ['pptx', 'ppt']:
                return cls.process_pptx(file.read())
            elif file_type == 'txt':
                return cls.process_txt(file.read())
            else:
                return "Unsupported file format"
